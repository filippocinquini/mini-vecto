"""
import_engine_map.py - turns a measured engine map (Excel from the test bench) into the CSV read by powertrain.py

  python3 import_engine_map.py "path/to/Reference 13-3636 LH08 and PLM_stripped.xlsx"

Writes data/engine_map_measured.csv (one row per measured operating point).
The measured map in the repository was provided by AB Volvo for the TRA545 course project.
Expected sheet layout: row 1 signal names (Spd_Engine, Trq_Engine, BSFC, Flw_Fuel, ...), row 2 descriptions,
row 3 units, data from row 4.
"""
import csv, os, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "data", "engine_map_measured.csv")
# signal in the Excel file -> column in the CSV
KEEP = [("Spd_Engine", "n_rpm"), ("Trq_Engine", "T_Nm"), ("Flw_Fuel", "fuel_g_per_s"), ("BSFC", "bsfc_g_per_kWh"),
        ("Spec_NOx", "nox_engine_out_g_per_kWh"), ("Flw_NOx", "nox_engine_out_mg_per_s"), ("Flw_CO2", "co2_g_per_s"),
        ("Flw_CO", "co_mg_per_s"), ("Flw_HC", "hc_mg_per_s"), ("Flw_Exhaust", "exhaust_kg_per_s"),
        ("Pwr_ExhaustHeat", "exhaust_heat_kW"), ("Pwr_CoolantHeat", "coolant_heat_kW")]

if len(sys.argv) < 2:
    sys.exit(__doc__)
ws = load_workbook(sys.argv[1], data_only=True).worksheets[0]
rows = list(ws.iter_rows(values_only=True))
col = {name: i for i, name in enumerate(rows[0]) if name}
missing = [a for a, _ in KEEP[:3] if a not in col]
if missing:
    sys.exit("Columns not found in the sheet: " + ", ".join(missing))
keep = [(a, b) for a, b in KEEP if a in col]
out = []
for r in rows[3:]:
    if r[col["Spd_Engine"]] is None:
        continue
    d = {b: r[col[a]] for a, b in keep}
    d["T_Nm"] = max(d["T_Nm"], 0.0)                      # the zero-torque points read -0.1 ... 0.1 Nm
    if d["T_Nm"] < 50:                                   # no brake power: values per kWh are meaningless
        for k in d:
            if k.endswith("_per_kWh"):
                d[k] = ""
    out.append({k: (round(v, 4) if isinstance(v, float) else v) for k, v in d.items()})
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", newline="") as f:
    wr = csv.DictWriter(f, fieldnames=[b for _, b in keep])
    wr.writeheader(); wr.writerows(out)
print(f"{len(out)} operating points written to {os.path.relpath(OUT, HERE)}")

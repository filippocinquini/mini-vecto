"""
export_excel.py - one Excel workbook per powertrain in results/

  results/consumption_Diesel.xlsx, _HVO100.xlsx, _H2-ICE.xlsx, _FCEV.xlsx, _BEV.xlsx
Each workbook: "Summary" (mean consumption per mission, live formulas on the step sheets),
               one sheet per mission (one row per 10 m step), "Notes" (model, parameters, sources).
Run after powertrain.py and electric.py:   python3 export_excel.py     (needs openpyxl)
"""
import csv, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "results")
CFG = json.load(open(os.path.join(HERE, "vehicles.json")))
EL = CFG["electric"]
MISSIONS = ["timber_empty", "timber_loaded", "urban_bus", "coach"]
TITLES = {"timber_empty": "Timber truck, empty leg (20 t)", "timber_loaded": "Timber truck, loaded leg (74 t)",
          "urban_bus": "Urban bus (14.7 t)", "coach": "Coach (17.2 t)"}
BOLD, BLUE = Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10, color="0000FF")
HEAD = PatternFill("solid", fgColor="D9E1F2")
BASE = ["s_m", "t_s", "dt_drive_s", "dwell_s", "v_kmh", "a_mps2", "grade_pct", "surface", "F_wheel_N", "P_wheel_kW",
        "E_wheel_traction_kWh", "E_wheel_braking_kWh"]
ICE = ["gear", "clutch_slip", "n_rpm", "T_engine_Nm", "P_engine_kW", "P_aux_kW", "P_driveline_loss_kW", "eta_engine",
       "E_fuel_MJ", "E_fuel_idle_MJ"]
UNITS = {"s_m": "position at step start [m]", "t_s": "time at step start [s]", "dt_drive_s": "driving time of the step [s]",
         "dwell_s": "standstill time at this point [s]", "v_kmh": "mean speed [km/h]", "a_mps2": "acceleration [m/s2]",
         "grade_pct": "road grade [%]", "surface": "road surface", "F_wheel_N": "wheel force [N]", "P_wheel_kW": "wheel power [kW]",
         "E_wheel_traction_kWh": "traction work at the wheels [kWh]", "E_wheel_braking_kWh": "braking work at the wheels [kWh]",
         "gear": "gear", "clutch_slip": "1 = launch with slipping clutch", "n_rpm": "engine speed [rpm]",
         "T_engine_Nm": "engine torque [Nm]", "P_engine_kW": "engine power [kW]", "P_aux_kW": "auxiliaries [kW]",
         "P_driveline_loss_kW": "gearbox + axle loss [kW]", "eta_engine": "engine brake efficiency [-]",
         "E_fuel_MJ": "fuel energy in the step, incl. idle [MJ]", "E_fuel_idle_MJ": "of which idling [MJ]",
         "diesel_g": "diesel in the step [g]", "diesel_cum_kg": "cumulative diesel [kg]",
         "hvo100_g": "HVO100 in the step [g]", "hvo100_cum_kg": "cumulative HVO100 [kg]",
         "h2_g": "hydrogen in the step [g]", "pilot_hvo_g": "pilot HVO100 in the step [g]",
         "pilot_share": "pilot share of fuel energy [-]", "h2_cum_kg": "cumulative hydrogen [kg]",
         "motor_torque_frac": "motor torque / max torque [-]", "motor_speed_frac": "motor speed / base speed [-]",
         "P_motor_el_kW": "motor electric power at the DC bus, negative = regeneration [kW]",
         "E_motor_el_kWh": "electricity to the motor [kWh]", "E_regen_el_kWh": "electricity regenerated [kWh]",
         "E_friction_brake_kWh": "braking work left to the friction brakes [kWh]", "E_aux_kWh": "auxiliaries [kWh]",
         "P_battery_kW": "battery power, + = discharge [kW]", "E_battery_kWh": "battery energy at the terminals, + = discharge [kWh]",
         "E_battery_loss_kWh": "battery internal loss [kWh]", "SOC": "battery state of charge at step end [-]",
         "E_battery_cum_kWh": "cumulative battery energy at the terminals [kWh]", "power_deficit": "1 = demand above the powertrain limit",
         "P_fc_net_kW": "fuel cell net power [kW]", "fc_load_frac": "fuel cell load / rated [-]",
         "eta_fc_sys_BOL": "fuel cell system efficiency, new [-]", "E_fc_net_kWh": "fuel cell net energy [kWh]",
         "E_dcdc_out_kWh": "energy after the DC/DC converter [kWh]", "h2_g_BOL": "hydrogen in the step, new fuel cell [g]",
         "h2_g_EOL": "hydrogen in the step, aged fuel cell [g]", "h2_cum_kg_BOL": "cumulative hydrogen, new fuel cell [kg]"}


def read(name):
    rows = list(csv.DictReader(open(os.path.join(HERE, name))))
    for r in rows:
        for k, v in r.items():
            if k != "surface":
                r[k] = float(v)
    return rows


def ice_rows(tag, fuel):
    rows = read(f"consumption_{tag}.csv")
    cum = 0.0
    for r in rows:
        r["E_wheel_traction_kWh"] = round(max(r["F_wheel_N"], 0) * 10 / 3.6e6, 6)
        r["E_wheel_braking_kWh"] = round(max(-r["F_wheel_N"], 0) * 10 / 3.6e6, 6)
        if fuel == "h2ice":
            r["h2_g"], r["pilot_hvo_g"] = r["h2ice_h2_g"], r["h2ice_pilot_hvo_g"]
            cum += r["h2_g"] / 1e3; r["h2_cum_kg"] = round(cum, 5)
        else:
            cum += r[fuel + "_g"] / 1e3; r[fuel + "_cum_kg"] = round(cum, 4)
    return rows


def data_sheet(wb, tag, cols, rows):
    ws = wb.create_sheet(tag)
    ws.append(cols); ws.append([UNITS.get(c, "") for c in cols])
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).font = BOLD; ws.cell(1, c).fill = HEAD
        ws.cell(2, c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[L(c)].width = 15
    ws.row_dimensions[2].height = 54
    for r in rows:
        ws.append([r[c] for c in cols])
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A1:{L(len(cols))}{len(rows) + 2}"
    return len(rows) + 2                                    # last data row


def build(name, title, cols, loader, inputs, ext, der, notes):
    """inputs: [(key, label, unit, value or {mission: value}, source note)]
       ext   : [(key, label, unit, fn(R) -> formula)] totals taken from the step sheets; R(col) = range of that column
       der   : [(key, label, unit, fn(c) -> formula, numfmt)] computed from rows above; c(key) = cell of that row"""
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name="Arial", size=10)
    ws = wb.active; ws.title = "Summary"
    last = {}
    for tag in MISSIONS:
        last[tag] = data_sheet(wb, tag, cols, loader(tag))
    ws["A1"] = title; ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = ("Mean consumption per mission. Every number below is a formula on the mission sheets (one row per 10 m step). "
                "Blue cells are inputs taken from vehicles.json.")
    hdr = ["", "unit"] + [TITLES[t] for t in MISSIONS] + ["Timber round trip (empty + loaded)", "note"]
    ws.append([]); ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        ws.cell(4, c).font = BOLD; ws.cell(4, c).fill = HEAD; ws.cell(4, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 42
    rowof, r = {}, 5

    def section(text):
        nonlocal r
        ws.cell(r, 1, text).font = BOLD; r += 1

    section("Inputs")
    for key, label, unit, val, src in inputs:
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for j, tag in enumerate(MISSIONS):
            v = val[tag] if isinstance(val, dict) else val
            ws.cell(r, 3 + j, v).font = BLUE
        if isinstance(val, dict):
            if val[MISSIONS[0]] == val[MISSIONS[1]]:
                ws.cell(r, 7, val[MISSIONS[0]]).font = BLUE
        else:
            ws.cell(r, 7, val).font = BLUE
        ws.cell(r, 8, src); rowof[key] = r; r += 1
    section("Totals over the mission")
    for key, label, unit, fn in ext:
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for j, tag in enumerate(MISSIONS):
            R = lambda col, tag=tag: f"'{tag}'!{L(cols.index(col) + 1)}3:{L(cols.index(col) + 1)}{last[tag]}"
            R.last = lambda col, tag=tag: f"'{tag}'!{L(cols.index(col) + 1)}{last[tag]}"
            ws.cell(r, 3 + j, fn(R))
        if key != "soc_end":
            ws.cell(r, 7, f"=C{r}+D{r}")
        rowof[key] = r; r += 1
    section("Mean consumption and efficiency")
    for key, label, unit, fn, fmt in der:
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for j in range(5):
            col = L(3 + j)
            c = lambda k, col=col: f"{col}{rowof[k]}"
            cell = ws.cell(r, 3 + j, fn(c)); cell.number_format = fmt
            if key in HEADLINE:
                cell.font = BOLD
        if key in HEADLINE:
            ws.cell(r, 1).font = BOLD
        rowof[key] = r; r += 1
    for rr in range(5, r):
        for cc in range(3, 8):
            if ws.cell(rr, cc).number_format == "General":
                ws.cell(rr, cc).number_format = "0.000"
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 12; ws.column_dimensions["H"].width = 90
    for cc in "CDEFG":
        ws.column_dimensions[cc].width = 19
    ws.freeze_panes = "C5"
    wn = wb.create_sheet("Notes", 1)
    wn.column_dimensions["A"].width = 150
    for line in notes:
        wn.append([line])
        if line.isupper():
            wn.cell(wn.max_row, 1).font = BOLD
        wn.cell(wn.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, f"consumption_{name}.xlsx")
    wb.save(path)
    print("wrote", path)


HEADLINE = {"l100", "kwhkm", "h2_100", "h2_bol", "h2_eol", "h2_mid", "bat_km", "grid_km"}
F = CFG["fuels"]
COMMON_EXT = [("km", "Distance", "km", lambda R: f"=COUNT({R('s_m')})*0.01"),
              ("min", "Duration", "min", lambda R: f"=(SUM({R('dt_drive_s')})+SUM({R('dwell_s')}))/60"),
              ("trac", "Traction work at the wheels", "kWh", lambda R: f"=SUM({R('E_wheel_traction_kWh')})"),
              ("brk", "Braking work at the wheels", "kWh", lambda R: f"=SUM({R('E_wheel_braking_kWh')})")]
COMMON_DER = [("vmean", "Mean speed", "km/h", lambda c: f"={c('km')}/({c('min')}/60)", "0.0"),
              ("trac_km", "Traction work at the wheels", "kWh/km", lambda c: f"={c('trac')}/{c('km')}", "0.000"),
              ("brk_km", "Braking work at the wheels", "kWh/km", lambda c: f"={c('brk')}/{c('km')}", "0.000")]
GENERAL = ["MODEL", "Backward vehicle simulation on 100 km distance-based mission cycles (10 m steps), see README.md.",
           "Wheel force = inertia + grade + rolling resistance + aerodynamic drag, same masses and resistances for every powertrain.",
           "The mass penalty of batteries and hydrogen tanks is NOT included yet (same vehicle mass as the diesel case): assumption to revisit for the LCA functional unit (payload).",
           "S = sourced value, D = derived from sources, A = assumption, to be confirmed with manufacturer data."]


def ice_notes(extra):
    v = CFG["vehicles"]
    return GENERAL + ["", "ENGINE POWERTRAIN", "Cycle -> I-Shift 12-speed (direct drive) -> engine map -> fuel. Gear = feasible gear with the lowest fuel use.",
                      "Engine map: semi-empirical, 12.8 L Volvo D13K540, peak brake efficiency 50 % (assumption); buses use the same map scaled by power.",
                      "Stops: engine idles (no start-stop). Overrun: fuel cut-off."] + extra + \
        ["", "VEHICLE NOTES"] + [f"{k}: {v[k]['notes']}" for k in v]


def el_notes(kind):
    e, ve = EL, EL["vehicles"]
    n = GENERAL + ["", "ELECTRIC DRIVE", e["edrive"]["note"], "", "BATTERY", e["battery"]["note"]]
    if kind == "fcev":
        f = e["fuel_cell"]
        n += ["", "FUEL CELL SYSTEM", f["note"],
              "Efficiency curve used (load fraction: new / aged): " + "; ".join(
                  f"{int(a*100)} %: {b:.3f} / {c:.3f}" for a, b, c in zip(f["load_frac"], f["eta_sys_BOL"], f["eta_sys_EOL"])),
              "Energy management: fuel cell power = filtered DC-bus demand (60 s) + correction on the battery state of charge, between 10 % and 100 % of rated power, ramp 10 %/s; it switches off when the request falls below 5 %.",
              "Hydrogen is corrected for the difference in battery state of charge between start and end: correction = battery energy change x raw hydrogen / (DC/DC output energy x battery charge efficiency).",
              "Mid-life value = mean of new (BOL) and aged (EOL) fuel cell."]
    n += ["", "VEHICLE NOTES"] + [f"{k}: {ve[k]['note']}" for k in ve]
    n += ["", "VALIDATION BENCHMARKS (literature)",
          "Fuel cell buses, measured: 7.7 kg H2/100 km for 12 m buses, 8.9 for 18 m (JIVE 2 final results, 2025); JIVE site averages 6.1-9.1.",
          "Fuel cell trucks at 36-40 t: about 8 kg/100 km implied by manufacturer range statements (Hyundai XCIENT, Mercedes GenH2); no measured value at 60-74 t.",
          "Battery electric timber truck, planning values: 2.5 kWh/km loaded, 1.5 kWh/km empty (Skogforsk Arbetsrapport 1176-2023).",
          "Battery electric city bus, measured on Gothenburg route 55: 0.93 kWh/km propulsion + 0.28 kWh/km HVO heater (Boren 2020)."]
    return n


def main():
    # ------------------------------------------------ Diesel and HVO100
    for fuel, name, label in (("diesel", "Diesel", "Diesel (EN590)"), ("hvo100", "HVO100", "HVO100")):
        cols = BASE + ICE + [fuel + "_g", fuel + "_cum_kg"]
        build(name, f"mini-VECTO - consumption, {label} engine", cols, lambda t, fuel=fuel: ice_rows(t, fuel),
              [("lhv", "Lower heating value", "MJ/kg", F[fuel]["LHV_MJ_per_kg"], "vehicles.json, fuels"),
               ("rho", "Density", "kg/l", F[fuel]["density_kg_per_l"], "vehicles.json, fuels")],
              COMMON_EXT + [("efuel", "Fuel energy", "MJ", lambda R: f"=SUM({R('E_fuel_MJ')})"),
                            ("eidle", "of which idling at stops", "MJ", lambda R: f"=SUM({R('E_fuel_idle_MJ')})"),
                            ("kg", "Fuel mass", "kg", lambda R, fuel=fuel: f"=SUM({R(fuel + '_g')})/1000")],
              [("l100", "Mean fuel consumption", "l/100 km", lambda c: f"={c('kg')}/{c('rho')}/{c('km')}*100", "0.00"),
               ("kg100", "Mean fuel consumption", "kg/100 km", lambda c: f"={c('kg')}/{c('km')}*100", "0.00"),
               ("kwhkm", "Fuel energy", "kWh/km", lambda c: f"={c('efuel')}/3.6/{c('km')}", "0.000")] + COMMON_DER +
              [("ttw", "Tank-to-wheel efficiency", "-", lambda c: f"={c('trac')}/({c('efuel')}/3.6)", "0.0%"),
               ("idle", "Idle share of fuel", "-", lambda c: f"={c('eidle')}/{c('efuel')}", "0.0%")],
              ice_notes(["Same engine map and efficiency for diesel and HVO100 (assumption); only heating value and density differ."]))
    # ------------------------------------------------ H2-ICE
    cols = BASE + ICE + ["h2_g", "pilot_hvo_g", "pilot_share", "h2_cum_kg"]
    build("H2-ICE", "mini-VECTO - consumption, hydrogen engine (HPDI with HVO100 pilot)", cols, lambda t: ice_rows(t, "h2ice"),
          [("lhv", "Lower heating value, hydrogen", "MJ/kg", F["h2"]["LHV_MJ_per_kg"], "vehicles.json, fuels"),
           ("lhvp", "Lower heating value, pilot HVO100", "MJ/kg", F["hvo100"]["LHV_MJ_per_kg"], "vehicles.json, fuels"),
           ("rho", "Density, pilot HVO100", "kg/l", F["hvo100"]["density_kg_per_l"], "vehicles.json, fuels")],
          COMMON_EXT + [("efuel", "Fuel energy (hydrogen + pilot)", "MJ", lambda R: f"=SUM({R('E_fuel_MJ')})"),
                        ("h2", "Hydrogen", "kg", lambda R: f"=SUM({R('h2_g')})/1000"),
                        ("pil", "Pilot HVO100", "kg", lambda R: f"=SUM({R('pilot_hvo_g')})/1000")],
          [("h2_100", "Mean hydrogen consumption", "kg/100 km", lambda c: f"={c('h2')}/{c('km')}*100", "0.00"),
           ("pil100", "Mean pilot HVO100 consumption", "l/100 km", lambda c: f"={c('pil')}/{c('rho')}/{c('km')}*100", "0.00"),
           ("pilsh", "Pilot share of fuel energy", "-", lambda c: f"={c('pil')}*{c('lhvp')}/{c('efuel')}", "0.0%"),
           ("kwhkm", "Fuel energy", "kWh/km", lambda c: f"={c('efuel')}/3.6/{c('km')}", "0.000")] + COMMON_DER +
          [("ttw", "Tank-to-wheel efficiency", "-", lambda c: f"={c('trac')}/({c('efuel')}/3.6)", "0.0%")],
          ice_notes(["Same engine map as diesel (assumption; AVL reports 49.1-50.1 % brake efficiency for a hydrogen HPDI engine). " + CFG["hpdi"]["note"]]))
    # ------------------------------------------------ electric
    ve = EL["vehicles"]
    veh = {t: ve[CFG["missions"][t]["vehicle"]] for t in MISSIONS}
    el_ext = COMMON_EXT + [("mot", "Electricity to the motor", "kWh", lambda R: f"=SUM({R('E_motor_el_kWh')})"),
                           ("reg", "Electricity regenerated", "kWh", lambda R: f"=SUM({R('E_regen_el_kWh')})"),
                           ("aux", "Auxiliaries", "kWh", lambda R: f"=SUM({R('E_aux_kWh')})"),
                           ("def", "Steps with demand above the powertrain limit", "-", lambda R: f"=SUM({R('power_deficit')})")]
    el_der = COMMON_DER + [("regsh", "Regenerated electricity / braking work", "-", lambda c: f"=IF({c('brk')}>0,{c('reg')}/{c('brk')},0)", "0.0%")]
    EB = ["gear", "motor_torque_frac", "motor_speed_frac", "P_motor_el_kW", "E_motor_el_kWh", "E_regen_el_kWh", "E_friction_brake_kWh", "E_aux_kWh"]
    # BEV
    cols = BASE + EB + ["P_battery_kW", "E_battery_kWh", "E_battery_loss_kWh", "SOC", "E_battery_cum_kWh", "power_deficit"]
    b = EL["battery"]
    build("BEV", "mini-VECTO - consumption, battery electric", cols, lambda t: read(f"consumption_bev_{t}.csv"),
          [("ebat", "Battery, installed", "kWh", {t: veh[t]["bev"]["battery_kWh"] for t in MISSIONS}, "vehicles.json, electric (see Notes)"),
           ("soc0", "State of charge at start", "-", {t: veh[t]["bev"]["soc_start"] for t in MISSIONS}, "assumption"),
           ("chg", "Charger efficiency", "-", b["charger_efficiency"], "Skogforsk Arbetsrapport 1176-2023"),
           ("bchg", "Battery efficiency while charging", "-", round(1 - b["loss_at_1C"] * b["depot_charge_C_rate"], 4),
            "assumption: 3 % loss at 1C, charging at 0.6C")],
          el_ext + [("bt", "Battery energy at the terminals (net of regeneration)", "kWh", lambda R: f"=SUM({R('E_battery_kWh')})"),
                    ("bl", "Battery internal loss", "kWh", lambda R: f"=SUM({R('E_battery_loss_kWh')})"),
                    ("soc_end", "State of charge at the end", "-", lambda R: f"={R.last('SOC')}")],
          [("bat_km", "Mean consumption from the battery", "kWh/km", lambda c: f"=({c('bt')}+{c('bl')})/{c('km')}", "0.000"),
           ("grid_km", "Mean consumption from the grid (incl. charging losses)", "kWh/km",
            lambda c: f"=({c('bt')}+{c('bl')})/{c('km')}/({c('chg')}*{c('bchg')})", "0.000"),
           ("used", "Share of installed battery used on 100 km", "-", lambda c: f"=IF({c('ebat')}>0,({c('bt')}+{c('bl')})/{c('ebat')},\"\")", "0.0%")] + el_der +
          [("ttw", "Battery-to-wheel efficiency", "-", lambda c: f"={c('trac')}/({c('bt')}+{c('bl')}+{c('reg')})", "0.0%")],
          el_notes("bev"))
    # FCEV
    cols = BASE + EB + ["P_fc_net_kW", "fc_load_frac", "eta_fc_sys_BOL", "E_fc_net_kWh", "E_dcdc_out_kWh", "P_battery_kW",
                        "E_battery_kWh", "E_battery_loss_kWh", "SOC", "h2_g_BOL", "h2_g_EOL", "h2_cum_kg_BOL", "power_deficit"]
    f = EL["fuel_cell"]
    build("FCEV", "mini-VECTO - consumption, fuel cell electric", cols, lambda t: read(f"consumption_fcev_{t}.csv"),
          [("pfc", "Fuel cell system, rated net power", "kW", {t: veh[t]["fcev"]["fc_net_kW"] for t in MISSIONS}, "vehicles.json, electric (see Notes)"),
           ("ebat", "Buffer battery, installed", "kWh", {t: veh[t]["fcev"]["battery_kWh"] for t in MISSIONS}, "vehicles.json, electric (see Notes)"),
           ("soc0", "State of charge at start", "-", f["ems"]["soc_ref"], "assumption"),
           ("lhv", "Lower heating value, hydrogen", "kWh/kg", round(120 / 3.6, 4), "120 MJ/kg"),
           ("bchg", "Battery efficiency while charging", "-", 1 - b["loss_at_1C"], "assumption: 3 % loss at 1C"),
           ("purge", "Purge share of hydrogen (already inside the consumption)", "-", f["purge_share_of_h2"], "assumption")],
          el_ext + [("efc", "Fuel cell net energy", "kWh", lambda R: f"=SUM({R('E_fc_net_kWh')})"),
                    ("edc", "Energy after the DC/DC converter", "kWh", lambda R: f"=SUM({R('E_dcdc_out_kWh')})"),
                    ("h2b", "Hydrogen, new fuel cell (BOL), before battery correction", "kg", lambda R: f"=SUM({R('h2_g_BOL')})/1000"),
                    ("h2e", "Hydrogen, aged fuel cell (EOL), before battery correction", "kg", lambda R: f"=SUM({R('h2_g_EOL')})/1000"),
                    ("dE", "Battery energy used (start minus end)", "kWh", lambda R: f"=SUM({R('E_battery_kWh')})+SUM({R('E_battery_loss_kWh')})"),
                    ("soc_end", "State of charge at the end", "-", lambda R: f"={R.last('SOC')}")],
          [("h2_bol", "Mean hydrogen consumption, new fuel cell (BOL)", "kg/100 km",
            lambda c: f"=({c('h2b')}+{c('dE')}*{c('h2b')}/({c('edc')}*{c('bchg')}))/{c('km')}*100", "0.00"),
           ("h2_eol", "Mean hydrogen consumption, aged fuel cell (EOL)", "kg/100 km",
            lambda c: f"=({c('h2e')}+{c('dE')}*{c('h2e')}/({c('edc')}*{c('bchg')}))/{c('km')}*100", "0.00"),
           ("h2_mid", "Mean hydrogen consumption, mid-life (mean of BOL and EOL)", "kg/100 km",
            lambda c: f"=AVERAGE({c('h2_bol')},{c('h2_eol')})", "0.00"),
           ("kwhkm", "Hydrogen energy, mid-life", "kWh/km", lambda c: f"={c('h2_mid')}/100*{c('lhv')}", "0.000"),
           ("etab", "Mean fuel cell system efficiency, BOL", "-", lambda c: f"={c('efc')}/({c('h2b')}*{c('lhv')})", "0.0%"),
           ("etae", "Mean fuel cell system efficiency, EOL", "-", lambda c: f"={c('efc')}/({c('h2e')}*{c('lhv')})", "0.0%"),
           ("pmean", "Mean fuel cell load", "-", lambda c: f"=IF({c('pfc')}>0,{c('efc')}/({c('min')}/60)/{c('pfc')},\"\")", "0.0%"),
           ("purg", "Hydrogen purged, mid-life", "g/km", lambda c: f"={c('h2_mid')}*10*{c('purge')}", "0.000")] + el_der +
          [("ttw", "Tank-to-wheel efficiency, mid-life", "-", lambda c: f"={c('trac')}/({c('h2_mid')}/100*{c('km')}*{c('lhv')})", "0.0%")],
          el_notes("fcev"))


if __name__ == "__main__":
    main()
